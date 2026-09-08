"""Synthetic-fixture tests for the deterministic Module 5A foundation."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

import fitz
from docx import Document
import yaml

from backend.app.services.rag.models import ExtractionStatus, SourceClass
from backend.app.services.rag.service import build_corpus, discover_documents, validate_corpus


def _write_fixture_config(root: Path) -> None:
    config = root / "config"
    config.mkdir(parents=True)
    payload = {
        "version": 1,
        "roots": [{"path": "data/rag", "allowed_extensions": [".pdf", ".docx", ".jsonl", ".json", ".zip"]}],
        "path_rules": [
            {
                "prefix": "data/rag/archive",
                "source_class": "ARCHIVE",
                "document_status": "UNKNOWN",
                "retrieval_eligible": False,
                "jurisdiction": "UNKNOWN",
                "applicable_workflows": [],
                "applicable_domains": ["UNKNOWN"],
            },
            {
                "prefix": "data/rag/restricted",
                "source_class": "RESTRICTED",
                "document_status": "UNKNOWN",
                "retrieval_eligible": False,
                "jurisdiction": "UNKNOWN",
                "applicable_workflows": [],
                "applicable_domains": ["UNKNOWN"],
                "restricted": True,
                "confidential": True,
            },
            {
                "prefix": "data/rag/governance",
                "source_class": "GOVERNANCE",
                "document_status": "SUPPORTING",
                "retrieval_eligible": False,
                "jurisdiction": "UNKNOWN",
                "applicable_workflows": ["POLICY_REGULATORY_INTELLIGENCE"],
                "applicable_domains": ["RESPONSIBLE_AI"],
            },
            {
                "prefix": "data/rag/primary/ireland/cru/historical_proposed",
                "source_class": "PRIMARY",
                "document_status": "PROPOSED",
                "retrieval_eligible": True,
                "jurisdiction": "IRELAND",
                "applicable_workflows": ["POLICY_REGULATORY_INTELLIGENCE"],
                "applicable_domains": ["GRID"],
            },
            {
                "prefix": "data/rag/primary/ireland",
                "source_class": "PRIMARY",
                "document_status": "CURRENT",
                "retrieval_eligible": True,
                "jurisdiction": "IRELAND",
                "applicable_workflows": ["SITE_FEASIBILITY"],
                "applicable_domains": ["GENERAL"],
            },
            {
                "prefix": "data/rag/supporting/industry",
                "source_class": "SUPPORTING_INDUSTRY",
                "document_status": "SUPPORTING",
                "retrieval_eligible": True,
                "jurisdiction": "IRELAND",
                "applicable_workflows": ["SITE_FEASIBILITY"],
                "applicable_domains": ["GENERAL"],
            },
            {
                "prefix": "data/rag/curated",
                "source_class": "CURATED",
                "document_status": "SUPPORTING",
                "retrieval_eligible": True,
                "jurisdiction": "IRELAND",
                "applicable_workflows": ["POLICY_REGULATORY_INTELLIGENCE"],
                "applicable_domains": ["DATA_CENTRE_POLICY"],
            },
        ],
    }
    (config / "rag_sources.yaml").write_text(yaml.safe_dump(payload, sort_keys=False), encoding="utf-8")
    (config / "rag_metadata_overrides.yaml").write_text("version: 1\ndocuments: {}\n", encoding="utf-8")


def _write_pdf(path: Path, pages: list[str | None]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf = fitz.open()
    for text in pages:
        page = pdf.new_page()
        if text:
            page.insert_text((72, 72), text)
    pdf.save(path)
    pdf.close()


def _write_docx(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("Feasibility heading", level=1)
    document.add_paragraph("A deterministic paragraph with provenance.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Requirement"
    table.rows[0].cells[1].text = "Evidence"
    document.save(path)


def _make_project(tmp_path: Path) -> Path:
    root = tmp_path / "project"
    root.mkdir()
    _write_fixture_config(root)
    _write_pdf(root / "data/rag/primary/ireland/current.pdf", ["CURRENT POLICY\nA current rule.", "Second page text."])
    (root / "data/rag/primary/ireland/duplicate.pdf").write_bytes((root / "data/rag/primary/ireland/current.pdf").read_bytes())
    _write_pdf(root / "data/rag/primary/ireland/cru/historical_proposed/proposed.pdf", ["PROPOSED DECISION\nHistorical proposal."])
    _write_pdf(root / "data/rag/restricted/confidential.pdf", ["CONFIDENTIAL TEXT MUST NOT BE CHUNKED."])
    _write_pdf(root / "data/rag/governance/responsible_ai/governance.pdf", ["Governance document."])
    _write_pdf(root / "data/rag/supporting/industry/submission.pdf", ["Industry perspective."])
    _write_docx(root / "data/rag/primary/ireland/policy.docx")
    (root / "data/rag/archive/original_packages/package.zip").parent.mkdir(parents=True, exist_ok=True)
    (root / "data/rag/archive/original_packages/package.zip").write_bytes(b"archive bytes")
    curated = root / "data/rag/curated/ireland_datacentre_screening"
    curated.mkdir(parents=True)
    record = {
        "id": "IE-DC-TEST-001",
        "title": {"en": "Curated test record"},
        "sources": [{"source_id": "SRC-TEST", "publisher": "Test", "title": "Test", "url": "https://example.com", "publication_date": None, "source_tier": "primary"}],
        "quality": {"review_required": False, "evidence_level": "primary_official", "confidence": "high"},
        "content": {"en": "Atomic curated content", "retrieval_text": "Atomic curated content"},
    }
    flagged = {
        "id": "IE-DC-TEST-002",
        "title": {"en": "Flagged curated record"},
        "sources": [{"source_id": "SRC-FLAGGED", "publisher": "Test", "title": "Test", "url": "https://example.com", "publication_date": None, "source_tier": "user_provided"}],
        "quality": {"review_required": False, "evidence_level": "user_provided_unverified", "confidence": "low"},
        "content": {"en": "Unsupported claim", "retrieval_text": "Unsupported claim"},
    }
    (curated / "rag_records.jsonl").write_text("\n".join(json.dumps(item) for item in (record, flagged)) + "\n", encoding="utf-8")
    (curated / "review_flags.jsonl").write_text(json.dumps({"flag_id": "FLAG-TEST", "source_id": "SRC-FLAGGED", "original_claim": "Unsupported claim"}) + "\n", encoding="utf-8")
    (curated / "source_registry.json").write_text("{}", encoding="utf-8")
    (curated / "rag_record_schema.json").write_text("{}", encoding="utf-8")
    (curated / "retrieval_eval_cases.jsonl").write_text(json.dumps({"case_id": "CASE-1"}) + "\n", encoding="utf-8")
    return root


def _registry(root: Path) -> list[dict]:
    return json.loads((root / "data/rag_processed/document_registry.json").read_text(encoding="utf-8"))["documents"]


def test_pdf_pages_and_chunks_preserve_provenance(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    report = build_corpus(root)
    registry = _registry(root)
    current = next(item for item in registry if item["source_path"].endswith("current.pdf"))
    chunks = [json.loads(line) for line in (root / "data/rag_processed/chunks.jsonl").read_text(encoding="utf-8").splitlines()]
    current_chunks = [chunk for chunk in chunks if chunk["document_id"] == current["document_id"]]
    assert current["page_count"] == 2
    assert current["extraction_status"] == ExtractionStatus.EXTRACTED.value
    assert current_chunks
    assert {chunk["page_start"] for chunk in current_chunks} == {1, 2}
    assert all(chunk["source_reference"]["document_id"] == current["document_id"] for chunk in current_chunks)
    assert all(chunk["source_reference"]["page_start"] == chunk["page_start"] for chunk in current_chunks)
    assert report["counts"]["total_active_chunks"] == len(chunks)


def test_docx_provenance_does_not_fabricate_pages(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    build_corpus(root)
    document = next(item for item in _registry(root) if item["filename"] == "policy.docx")
    chunks = [json.loads(line) for line in (root / "data/rag_processed/chunks.jsonl").read_text(encoding="utf-8").splitlines()]
    docx_chunks = [chunk for chunk in chunks if chunk["document_id"] == document["document_id"]]
    assert document["page_count"] is None
    assert docx_chunks
    assert all(chunk["page_start"] is None and chunk["page_end"] is None for chunk in docx_chunks)
    assert any(chunk["source_reference"]["paragraph_start"] is not None for chunk in docx_chunks)


def test_authority_status_and_exclusion_rules(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    build_corpus(root)
    registry = _registry(root)
    current = next(item for item in registry if item["filename"] == "current.pdf")
    proposed = next(item for item in registry if item["filename"] == "proposed.pdf")
    archive = next(item for item in registry if item["filename"] == "package.zip")
    restricted = next(item for item in registry if item["filename"] == "confidential.pdf")
    governance = next(item for item in registry if item["filename"] == "governance.pdf")
    industry = next(item for item in registry if item["filename"] == "submission.pdf")
    assert current["source_class"] == SourceClass.PRIMARY.value and current["retrieval_eligible"] is True
    assert proposed["document_status"] == "PROPOSED" and proposed["requires_human_review"] is False
    assert archive["retrieval_eligible"] is False and archive["extraction_status"] == "SKIPPED_ARCHIVE"
    assert restricted["restricted"] is True and restricted["chunk_count"] == 0
    assert governance["retrieval_eligible"] is False and governance["chunk_count"] == 0
    assert industry["source_class"] == SourceClass.SUPPORTING_INDUSTRY.value
    assert industry["source_class"] != SourceClass.PRIMARY.value


def test_duplicate_hashes_create_one_active_chunk_set(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    build_corpus(root)
    registry = _registry(root)
    duplicates = [item for item in registry if item["filename"] in {"current.pdf", "duplicate.pdf"}]
    chunks = [json.loads(line) for line in (root / "data/rag_processed/chunks.jsonl").read_text(encoding="utf-8").splitlines()]
    assert len({item["sha256"] for item in duplicates}) == 1
    assert sum(item["duplicate_of_document_id"] is not None for item in duplicates) == 1
    assert len({chunk["document_id"] for chunk in chunks if chunk["document_id"] in {item["document_id"] for item in duplicates}}) == 1


def test_curated_records_are_atomic_and_review_flags_attach(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    build_corpus(root)
    records = [json.loads(line) for line in (root / "data/rag_processed/curated_records.jsonl").read_text(encoding="utf-8").splitlines()]
    registry = _registry(root)
    curated_document = next(item for item in registry if item["filename"] == "rag_records.jsonl")
    assert len(records) == 2
    assert all(item["document_id"] == curated_document["document_id"] for item in records)
    assert any(item["record_id"] == "IE-DC-TEST-002" and item["review_flags"] for item in records)
    assert any(item["record_id"] == "IE-DC-TEST-002" and item["retrieval_eligible"] is False for item in records)
    assert all(item["chunk_count"] == 0 for item in registry if item["filename"] == "rag_records.jsonl")
    source_registry_document = next(item for item in registry if item["filename"] == "source_registry.json")
    chunks = [json.loads(line) for line in (root / "data/rag_processed/chunks.jsonl").read_text(encoding="utf-8").splitlines()]
    assert all(chunk["document_id"] != source_registry_document["document_id"] for chunk in chunks)


def test_unknown_metadata_and_bad_pdf_do_not_stop_other_sources(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    bad = root / "data/rag/primary/ireland/bad.pdf"
    bad.write_bytes(b"not a PDF")
    before_hash = hashlib.sha256((root / "data/rag/primary/ireland/current.pdf").read_bytes()).hexdigest()
    build_corpus(root)
    registry = _registry(root)
    bad_document = next(item for item in registry if item["filename"] == "bad.pdf")
    current = next(item for item in registry if item["filename"] == "current.pdf")
    assert bad_document["extraction_status"] == ExtractionStatus.FAILED.value
    assert current["extraction_status"] == ExtractionStatus.EXTRACTED.value
    assert current["publication_date"] is None
    assert hashlib.sha256((root / "data/rag/primary/ireland/current.pdf").read_bytes()).hexdigest() == before_hash


def test_validate_checks_registry_chunks_and_source_hashes(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    build_corpus(root)
    result = validate_corpus(root)
    assert result["valid"] is True
    assert result["documents"] > 0
    assert result["chunks"] > 0


def test_pdf_blank_page_is_recorded_without_ocr(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    _write_pdf(root / "data/rag/primary/ireland/blank-page.pdf", ["Text page", None])
    build_corpus(root)
    document = next(item for item in _registry(root) if item["filename"] == "blank-page.pdf")
    assert document["no_text_pages"] == [2]
    assert document["extraction_status"] == ExtractionStatus.EXTRACTED.value


def test_pdf_citation_contains_exact_page_and_source_path(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    build_corpus(root)
    current = next(item for item in _registry(root) if item["filename"] == "current.pdf")
    chunks = [json.loads(line) for line in (root / "data/rag_processed/chunks.jsonl").read_text(encoding="utf-8").splitlines()]
    chunk = next(item for item in chunks if item["document_id"] == current["document_id"])
    assert chunk["source_reference"]["source_path"] == current["source_path"]
    assert chunk["source_reference"]["page_start"] == chunk["page_start"]
    assert chunk["source_reference"]["locator"].startswith("page ")


def test_primary_current_policy_is_retrieval_eligible(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    build_corpus(root)
    current = next(item for item in _registry(root) if item["filename"] == "current.pdf")
    assert current["source_class"] == "PRIMARY"
    assert current["document_status"] == "CURRENT"
    assert current["retrieval_eligible"] is True


def test_proposed_policy_does_not_become_current(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    build_corpus(root)
    proposed = next(item for item in _registry(root) if item["filename"] == "proposed.pdf")
    assert proposed["document_status"] == "PROPOSED"
    assert proposed["document_status"] != "CURRENT"


def test_archive_document_is_registered_but_not_chunked(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    build_corpus(root)
    document = next(item for item in _registry(root) if item["filename"] == "package.zip")
    assert document["source_class"] == "ARCHIVE"
    assert document["retrieval_eligible"] is False
    assert document["chunk_count"] == 0


def test_restricted_document_is_registered_but_not_chunked(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    build_corpus(root)
    document = next(item for item in _registry(root) if item["filename"] == "confidential.pdf")
    assert document["restricted"] is True
    assert document["extraction_status"] == ExtractionStatus.SKIPPED_RESTRICTED.value
    assert document["chunk_count"] == 0


def test_governance_document_is_excluded_from_normal_corpus(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    build_corpus(root)
    document = next(item for item in _registry(root) if item["filename"] == "governance.pdf")
    assert document["source_class"] == "GOVERNANCE"
    assert document["retrieval_eligible"] is False
    assert document["extraction_status"] == ExtractionStatus.SKIPPED_GOVERNANCE.value


def test_industry_document_is_supporting_only(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    build_corpus(root)
    document = next(item for item in _registry(root) if item["filename"] == "submission.pdf")
    chunks = [json.loads(line) for line in (root / "data/rag_processed/chunks.jsonl").read_text(encoding="utf-8").splitlines()]
    assert document["source_class"] == "SUPPORTING_INDUSTRY"
    assert all(chunk["source_class"] != "PRIMARY" for chunk in chunks if chunk["document_id"] == document["document_id"])


def test_sha256_duplicate_relationship_is_recorded(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    build_corpus(root)
    documents = [item for item in _registry(root) if item["filename"] in {"current.pdf", "duplicate.pdf"}]
    assert documents[0]["sha256"] == documents[1]["sha256"]
    assert sum(item["duplicate_of_document_id"] is not None for item in documents) == 1


def test_duplicate_source_does_not_duplicate_active_chunks(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    build_corpus(root)
    documents = [item for item in _registry(root) if item["filename"] in {"current.pdf", "duplicate.pdf"}]
    chunks = [json.loads(line) for line in (root / "data/rag_processed/chunks.jsonl").read_text(encoding="utf-8").splitlines()]
    ids = {document["document_id"] for document in documents}
    assert len({chunk["document_id"] for chunk in chunks if chunk["document_id"] in ids}) == 1


def test_curated_record_is_not_rechunked(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    build_corpus(root)
    document = next(item for item in _registry(root) if item["filename"] == "rag_records.jsonl")
    assert document["extraction_status"] == ExtractionStatus.CURATED_ATOMIC.value
    assert document["curated_record_count"] == 2
    assert document["chunk_count"] == 0


def test_review_flag_is_attached_and_disables_flagged_record(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    build_corpus(root)
    records = [json.loads(line) for line in (root / "data/rag_processed/curated_records.jsonl").read_text(encoding="utf-8").splitlines()]
    flagged = next(item for item in records if item["record_id"] == "IE-DC-TEST-002")
    assert flagged["review_flag_status"] == "MATCHED"
    assert flagged["review_flags"]
    assert flagged["retrieval_eligible"] is False


def test_review_flag_audit_separates_matched_and_unmatched_claims(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    flags_path = root / "data/rag/curated/ireland_datacentre_screening/review_flags.jsonl"
    with flags_path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps({"flag_id": "FLAG-UNMATCHED", "source_id": "SRC-NOT-IN-RECORDS", "original_claim": "Claim absent from cleaned records"}) + "\n")

    build_corpus(root)
    audit = json.loads((root / "data/rag_processed/review_flag_audit.json").read_text(encoding="utf-8"))
    by_id = {item["flag_id"]: item for item in audit["flags"]}

    assert by_id["FLAG-TEST"]["assessment"] == "B_ACTIVE_CURATED_RECORD_MATCH"
    assert by_id["FLAG-TEST"]["trust_treatment"] == "REVIEW_REQUIRED_MATCHED"
    assert by_id["FLAG-UNMATCHED"]["assessment"] == "A_ORIGINAL_SOURCE_CLAIM_ABSENT_FROM_CLEANED_RECORDS"
    assert by_id["FLAG-UNMATCHED"]["matched_record_ids"] == []
    assert by_id["FLAG-UNMATCHED"]["trust_treatment"] == "NOT_TRUSTED_UNMATCHED"


def test_primary_quality_and_no_text_reports_are_complete(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    _write_pdf(root / "data/rag/primary/ireland/visual.pdf", ["Title page", None])
    build_corpus(root)

    registry = _registry(root)
    primary_ids = {item["document_id"] for item in registry if item["source_class"] == SourceClass.PRIMARY.value}
    quality = json.loads((root / "data/rag_processed/primary_source_quality.json").read_text(encoding="utf-8"))
    no_text = json.loads((root / "data/rag_processed/no_text_pages.json").read_text(encoding="utf-8"))

    assert {item["document_id"] for item in quality["documents"]} == primary_ids
    visual_id = next(item["document_id"] for item in registry if item["filename"] == "visual.pdf")
    assert {item["document_id"] for item in no_text["pages"]} == {visual_id}
    assert no_text["pages"][0]["page"] == 2


def test_curated_metadata_and_evaluation_cases_are_not_knowledge_chunks(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    build_corpus(root)
    registry = _registry(root)
    chunks = [json.loads(line) for line in (root / "data/rag_processed/chunks.jsonl").read_text(encoding="utf-8").splitlines()]
    excluded_ids = {
        item["document_id"]
        for item in registry
        if item["filename"] in {"source_registry.json", "review_flags.jsonl", "rag_record_schema.json", "retrieval_eval_cases.jsonl"}
    }
    assert all(chunk["document_id"] not in excluded_ids for chunk in chunks)


def test_unknown_date_is_null_and_source_hash_is_unchanged(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    source = root / "data/rag/primary/ireland/current.pdf"
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    build_corpus(root)
    document = next(item for item in _registry(root) if item["filename"] == "current.pdf")
    assert document["publication_date"] is None
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before


def test_bad_pdf_does_not_prevent_other_documents_from_processing(tmp_path: Path) -> None:
    root = _make_project(tmp_path)
    (root / "data/rag/primary/ireland/bad.pdf").write_bytes(b"not a PDF")
    build_corpus(root)
    registry = _registry(root)
    bad = next(item for item in registry if item["filename"] == "bad.pdf")
    good = next(item for item in registry if item["filename"] == "current.pdf")
    assert bad["extraction_status"] == ExtractionStatus.FAILED.value
    assert good["extraction_status"] == ExtractionStatus.EXTRACTED.value
